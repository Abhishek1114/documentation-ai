"""Render a validated GMPRecord as a DOCX QA-review draft. Deterministic template: no model call."""
from docx import Document

from src.render.markdown import TITLES, fmt_num, fmt_opt, limit_check, review_flags

NOTE = "Draft for QA review. Generated from extracted data; not a validated GMP document."
KV_HDR = ["Field", "Value"]
STEPS_HDR = ["No", "Description", "Result", "Initials"]
MEAS_HDR = ["Parameter", "Value", "Unit", "Limit min", "Limit max", "Limit check"]
SIGN_HDR = ["Role", "Name", "Signature", "Date"]


def plain(s) -> str:
    return str(s).replace("\n", " ").strip()


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    return t


def render_docx(rec):
    doc = Document()
    doc.add_heading(TITLES[rec.document_type], 0)
    doc.add_paragraph().add_run(NOTE).italic = True

    doc.add_heading("Document details", 1)
    _table(doc, KV_HDR, [
        ["Document ID", plain(rec.document_id)],
        ["Document type", rec.document_type],
        ["Date", rec.date.isoformat()],
        ["Performed by", plain(rec.performed_by)],
        ["Reviewed by", plain(rec.reviewed_by) if rec.reviewed_by else "Not recorded"],
    ])

    doc.add_heading("Equipment", 1)
    _table(doc, KV_HDR, [
        ["ID", plain(rec.equipment.id)],
        ["Name", plain(rec.equipment.name)],
        ["Location", plain(rec.equipment.location)],
    ])

    doc.add_heading("Steps", 1)
    if rec.steps:
        _table(doc, STEPS_HDR, [[str(s.no), plain(s.description), s.result, plain(s.initials)]
                                for s in rec.steps])
    else:
        doc.add_paragraph("No steps recorded.")

    doc.add_heading("Measurements", 1)
    if rec.measurements:
        _table(doc, MEAS_HDR, [[plain(m.parameter), fmt_num(m.value), plain(m.unit), fmt_opt(m.limit_min),
                                fmt_opt(m.limit_max), limit_check(m)] for m in rec.measurements])
    else:
        doc.add_paragraph("No measurements recorded.")

    doc.add_heading("Deviations", 1)
    if rec.deviations:
        for d in rec.deviations:
            doc.add_paragraph(plain(d), style="List Bullet")
    else:
        doc.add_paragraph("None recorded.")

    doc.add_heading("Review flags", 1)
    flags = review_flags(rec)
    if flags:
        for f in flags:
            doc.add_paragraph(f.replace("\\|", "|"), style="List Bullet")  # review_flags escapes pipes for Markdown
    else:
        doc.add_paragraph("None.")

    doc.add_heading("Sign-off", 1)
    _table(doc, SIGN_HDR, [
        ["Performed by", plain(rec.performed_by), "", ""],
        ["Reviewed by", plain(rec.reviewed_by) if rec.reviewed_by else "", "", ""],
    ])
    return doc
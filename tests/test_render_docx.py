import io

from docx import Document

from src.render.docx_checks import check_docx
from src.render.docx_render import render_docx
from src.schema.models import GMPRecord

BASE = {
    "document_id": "CR-2026-00601", "document_type": "cleaning_record",
    "equipment": {"id": "MX-101", "name": "Ribbon Blender", "location": "Room B-12"},
    "date": "2026-07-02", "performed_by": "A. Patel", "reviewed_by": None,
    "steps": [
        {"no": 1, "description": "Rinse | until clear", "result": "pass", "initials": "AP"},
        {"no": 2, "description": "Wipe with 70% IPA", "result": "fail", "initials": "KS"},
    ],
    "measurements": [
        {"parameter": "TOC", "value": 612.5, "unit": "ppb", "limit_min": 0.0, "limit_max": 500.0},
        {"parameter": "pH of final rinse", "value": 7.3, "unit": "pH", "limit_min": None, "limit_max": None},
    ],
    "deviations": ["Residue found on gasket"],
}


def rec(**over):
    return GMPRecord(**{**BASE, **over})


def roundtrip(r):
    buf = io.BytesIO()
    render_docx(r).save(buf)      # serialise to bytes, then read back like a real file
    buf.seek(0)
    return Document(buf)


def meas_table(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text == "Parameter")


def test_roundtrip_passes_all_checks():
    r = rec()
    assert check_docx(roundtrip(r), r) == []


def test_pipe_text_is_kept_literally():
    r = rec()
    doc = roundtrip(r)
    steps = next(t for t in doc.tables if t.rows[0].cells[0].text == "No")
    assert steps.cell(1, 1).text == "Rinse | until clear"


def test_empty_lists_render_and_pass():
    r = rec(steps=[], measurements=[], deviations=[], reviewed_by="S. Das")
    assert check_docx(roundtrip(r), r) == []


def test_tampered_value_is_caught():
    r = rec()
    doc = roundtrip(r)
    meas_table(doc).cell(1, 1).text = "512.5"
    errs = check_docx(doc, r)
    assert any("differs from record" in e for e in errs) and any("512.5" in e for e in errs)


def test_invented_number_is_caught():
    r = rec()
    doc = roundtrip(r)
    doc.add_paragraph("Total: 99 items")
    assert any("99" in e for e in check_docx(doc, r))


def test_missing_section_is_caught():
    r = rec()
    doc = roundtrip(r)
    for p in doc.paragraphs:
        if p.text == "Deviations":
            p.text = "Notes"
    assert any("Deviations" in e for e in check_docx(doc, r))